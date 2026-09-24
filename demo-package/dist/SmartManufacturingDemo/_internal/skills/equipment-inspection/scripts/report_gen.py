# -*- coding: utf-8 -*-
"""
设备点检 Skill - 报告生成脚本
=============================
功能：基于分析结果生成点检日报、告警文件，并将结果同步到 Excel 台账。

用法示例：
    python report_gen.py                                   # 使用默认分析结果生成
    python report_gen.py --result ../output/analysis_result.json
    python report_gen.py --report-dir ../output --ledger ../output/inspection_ledger.xlsx

说明：
- 日报使用 templates/daily_report.md 渲染。
- 告警使用 templates/alert_template.md 渲染（按告警等级分别生成）。
- 台账使用 openpyxl 写入（依赖：openpyxl，已安装）。
- 真实部署中，如需发送企业微信告警，可在 `_send_wecom_alert` 中接入企业微信 MCP。

依赖：openpyxl（Excel台账）。
"""
import argparse
import json
import os
import sys
from datetime import datetime

SKILL_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.path.join(SKILL_ROOT, "templates")
OUTPUT_DIR = os.path.join(SKILL_ROOT, "output")
DAILY_TMPL = os.path.join(TEMPLATES_DIR, "daily_report.md")
ALERT_TMPL = os.path.join(TEMPLATES_DIR, "alert_template.md")


def _render(tmpl_path, **kwargs):
    with open(tmpl_path, "r", encoding="utf-8") as f:
        tpl = f.read()
    for k, v in kwargs.items():
        tpl = tpl.replace("{{" + k + "}}", str(v))
    return tpl


def _compose_alert_files(result, report_dir):
    """按告警等级分组生成告警文件，返回生成的列表。"""
    files = []
    alerts_by_status = {}
    for a in result["alerts"]:
        alerts_by_status.setdefault(a["status"], []).append(a)
    for status, alerts in alerts_by_status.items():
        parts = []
        for i, a in enumerate(alerts, 1):
            parts.append(_render(
                ALERT_TMPL,
                generated_at=result["generated_at"],
                alert_level=status,
                device_id=a["device_id"],
                device_name=a["device_name"],
                enterprise=a["enterprise"],
                param=a["param"],
                unit=a["unit"],
                current_value=a["value"],
                normal_range=a["normal_range"],
                critical_range=a["crit_range"],
                over_count=a["over_count"],
                action_hint=a["action_hint"],
            ))
        fname = os.path.join(report_dir, f"alert_{status}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md")
        with open(fname, "w", encoding="utf-8") as f:
            f.write("\n\n---\n\n".join(parts))
        files.append(fname)
    return files


def _gen_daily_report(result, report_dir):
    devices = result["devices"]
    rows = []
    for d in devices:
        for it in d["items"]:
            rows.append(
                f"| {d['device_id']} | {d['device_name']} | {d['enterprise']} | "
                f"{it['param']} | {it['unit']} | {it['value']} | {it['status']} | {it['note']} |"
            )
    s = result["summary"]
    summary_text = (
        f"当前全部 {sum(s.values())} 台设备中，正常 {s['正常']} 台；"
        f"预警 {s['预警']} 台、严重 {s['严重']} 台、紧急 {s['紧急']} 台。"
    )
    if result["alerts"]:
        summary_text += f"共触发 {len(result['alerts'])} 项告警，详见下方明细与告警文件。"
    else:
        summary_text += "本次点检全部正常，无需处置。"
    recs = []
    if s["紧急"]:
        recs.append("- **紧急项**：立即停机检修，禁止带病运行，同步通知设备工程师与安全管理。")
    if s["严重"]:
        recs.append("- **严重项**：当班处理，停止作业检查，必要时切换备用产线并通报班组长。")
    if s["预警"]:
        recs.append("- **预警项**：24小时内复核数据趋势，持续越界请联系设备工程师排查。")
    if not recs:
        recs.append("- 全部正常，维持现有点检计划即可。")
    recs.append("- 建议对连续出现 over_count 的点检项，纳入月度预防性维护计划。")

    html = _render(
        DAILY_TMPL,
        generated_at=result["generated_at"],
        scope_desc=f"全部设备（数据时间 {result.get('data_time', '未知')}，来源 {result.get('data_source', '未知')}）",
        data_source=result.get("data_source", "未知"),
        device_count=sum(s.values()),
        normal_count=s["正常"],
        warning_count=s["预警"] + s["严重"] + s["紧急"],
        critical_count=s["严重"] + s["紧急"],
        summary_text=summary_text,
        device_rows="\n".join(rows) if rows else "（无点检项）",
        recommendations="\n".join(recs),
    )
    fname = os.path.join(report_dir, f"daily_report_{datetime.now().strftime('%Y%m%d')}.md")
    with open(fname, "w", encoding="utf-8") as f:
        f.write(html)
    return fname


def _update_ledger(result, ledger_path):
    """将本次点检结果逐项追加到 Excel 台账。"""
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("[report_gen] 提示：未安装 openpyxl，跳过 Excel 台账同步。")
        return False

    os.makedirs(os.path.dirname(ledger_path), exist_ok=True)
    headers = ["日期", "设备代码", "设备名称", "所属企业", "点检参数", "单位",
               "当前值", "状态", "正常限", "24h超限次数", "说明"]
    if os.path.exists(ledger_path):
        wb = load_workbook(ledger_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "点检台账"
        ws.append(headers)

    today = datetime.now().strftime("%Y-%m-%d")
    for d in result["devices"]:
        for it in d["items"]:
            ws.append([
                today, d["device_id"], d["device_name"], d["enterprise"],
                it["param"], it["unit"], it["value"], it["status"],
                f"{it['normal_low']}~{it['normal_high']}", it["over_count_24h"], it["note"],
            ])
    wb.save(ledger_path)
    return True


def _send_wecom_alert(alert):
    """企业微信告警发送接口（占位）。

    真实部署时，在此接入企业微信 MCP，按告警等级推送：
        alert["status"] in {"预警", "严重", "紧急"}
    """
    # TODO: 接入企业微信 MCP 后启用
    #   await call 企业微信MCP.send_text(users=[...], content=...)
    pass


def main():
    parser = argparse.ArgumentParser(description="设备点检 - 报告生成")
    parser.add_argument("--result", default=os.path.join(OUTPUT_DIR, "analysis_result.json"),
                        help="分析结果JSON路径")
    parser.add_argument("--report-dir", default=OUTPUT_DIR, help="报告输出目录")
    parser.add_argument("--ledger", default=os.path.join(OUTPUT_DIR, "inspection_ledger.xlsx"),
                        help="Excel台账路径")
    args = parser.parse_args()

    if not os.path.exists(args.result):
        print("[report_gen] 错误：未找到分析结果，请先运行 data_fetch.py + analysis.py。")
        sys.exit(1)

    with open(args.result, "r", encoding="utf-8") as f:
        result = json.load(f)

    os.makedirs(args.report_dir, exist_ok=True)

    daily = _gen_daily_report(result, args.report_dir)
    print(f"[report_gen] 日报已生成 → {daily}")

    alert_files = _compose_alert_files(result, args.report_dir)
    for af in alert_files:
        print(f"[report_gen] 告警文件已生成 → {af}")

    # 企业微信告警（占位，接入 MCP 后生效）
    for a in result.get("alerts", []):
        if a["status"] in ("严重", "紧急"):
            _send_wecom_alert(a)

    ok = _update_ledger(result, args.ledger)
    if ok:
        print(f"[report_gen] Excel 台账已同步 → {args.ledger}")


if __name__ == "__main__":
    sys.exit(main())